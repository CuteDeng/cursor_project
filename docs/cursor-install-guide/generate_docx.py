#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""生成 Cursor Windows / macOS 注册到安装图文 Word 文档。"""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
IMG = ROOT / "images"
OUT = ROOT / "Cursor客户端注册与安装指南_Windows与Mac.docx"


def set_run_font(run, name="微软雅黑", size=11, bold=False, color=None):
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if color is not None:
        run.font.color.rgb = color


def add_para(doc, text, *, size=11, bold=False, color=None, align=None, space_after=8, space_before=0):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(space_before)
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color)
    return p


def add_heading_cn(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        set_run_font(run, size=18 if level == 1 else (14 if level == 2 else 12), bold=True)
    return h


def add_image(doc, path, *, width_cm=14.5, caption=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run()
    run.add_picture(str(path), width=Cm(width_cm))
    if caption:
        add_para(
            doc,
            caption,
            size=9,
            color=RGBColor(0x66, 0x66, 0x66),
            align=WD_ALIGN_PARAGRAPH.CENTER,
            space_after=12,
        )


def add_bullets(doc, items, *, size=11):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        run = p.add_run(item)
        set_run_font(run, size=size)


def add_numbered(doc, items, *, size=11):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        run = p.add_run(item)
        set_run_font(run, size=size)


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ""
        p = hdr[i].paragraphs[0]
        run = p.add_run(h)
        set_run_font(run, size=10, bold=True)
    for r_idx, row in enumerate(rows):
        cells = table.rows[r_idx + 1].cells
        for c_idx, val in enumerate(row):
            cells[c_idx].text = ""
            p = cells[c_idx].paragraphs[0]
            run = p.add_run(val)
            set_run_font(run, size=10)
    doc.add_paragraph()


def build():
    doc = Document()

    # 页面边距
    for section in doc.sections:
        section.top_margin = Cm(2.2)
        section.bottom_margin = Cm(2.2)
        section.left_margin = Cm(2.4)
        section.right_margin = Cm(2.4)

    # ========== 封面 ==========
    add_image(doc, IMG / "00-cover.png", width_cm=15.5)
    add_para(
        doc,
        "Cursor 客户端注册与安装指南",
        size=26,
        bold=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        space_before=12,
        space_after=6,
    )
    add_para(
        doc,
        "Windows & macOS 完整图文教程（从账号注册到首次使用）",
        size=14,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        space_after=6,
        color=RGBColor(0x44, 0x44, 0x44),
    )
    add_para(
        doc,
        "官方下载：https://cursor.com/download\n官方文档：https://cursor.com/help/getting-started/install",
        size=10,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        space_after=6,
        color=RGBColor(0x55, 0x55, 0x55),
    )
    add_para(
        doc,
        "文档版本：v1.0　｜　适用平台：Windows 10/11、macOS 12+",
        size=10,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        space_after=18,
        color=RGBColor(0x66, 0x66, 0x66),
    )

    # ========== 目录说明 ==========
    add_heading_cn(doc, "文档说明", 1)
    add_para(
        doc,
        "本文档面向首次使用 Cursor 的用户，按「注册账号 → 下载安装包 → 安装客户端 → 登录验证 → 打开项目」的顺序，分别说明 Windows 与 macOS 的完整操作步骤。配图为界面示意，实际界面可能随版本更新略有差异，请以官网与客户端当前界面为准。",
    )
    add_image(doc, IMG / "07-flow.png", width_cm=15, caption="图 0　整体流程：注册 → 下载 → 安装 → 登录 → 开始编码")

    add_heading_cn(doc, "一、开始之前：系统要求与准备事项", 1)
    add_heading_cn(doc, "1.1 系统要求", 2)
    add_table(
        doc,
        ["项目", "Windows", "macOS"],
        [
            ["操作系统", "Windows 10 及以上（建议 Windows 11）", "macOS 12 Monterey 及以上"],
            ["芯片架构", "x64 / ARM64", "Apple Silicon（M 系列）或 Intel"],
            ["安装包格式", ".exe 安装程序", ".dmg 磁盘镜像"],
            ["网络", "可访问 cursor.com 及相关认证域名", "同左"],
            ["账号", "需要 Cursor 账号（可免费注册 Hobby）", "同左"],
        ],
    )
    add_heading_cn(doc, "1.2 建议提前准备", 2)
    add_bullets(
        doc,
        [
            "可用的电子邮箱，或 Google / GitHub 账号（用于注册与登录）",
            "稳定的网络连接（下载安装包约 200MB 量级，登录需访问认证服务）",
            "管理员权限（部分 Windows 系统安装模式或公司电脑可能需要）",
            "若曾使用 VS Code，可在首次启动时一键导入扩展、快捷键与主题",
        ],
    )

    # ========== 注册 ==========
    add_heading_cn(doc, "二、注册 Cursor 账号", 1)
    add_para(
        doc,
        "官方建议：在安装或首次启动前先完成账号注册，这样安装后即可直接登录使用 AI 功能（补全、Chat、Agent 等）。Hobby（免费）计划即可体验核心能力。",
    )
    add_heading_cn(doc, "2.1 方式 A：官网注册（推荐）", 2)
    add_numbered(
        doc,
        [
            "用浏览器打开 https://cursor.com",
            "点击 Sign Up / 注册（或进入登录页后选择创建账号）",
            "选择注册方式：邮箱注册，或 Continue with Google / Continue with GitHub",
            "若使用邮箱：填写邮箱并按提示设置密码，查收验证邮件并完成验证",
            "注册成功后即可进入下载页，或稍后在客户端中登录",
        ],
    )
    add_image(doc, IMG / "02-signup.png", width_cm=14.5, caption="图 1　账号注册示意：邮箱 / Google / GitHub")

    add_heading_cn(doc, "2.2 方式 B：客户端内注册", 2)
    add_numbered(
        doc,
        [
            "先完成下文的下载与安装，并启动 Cursor",
            "在欢迎页点击 Sign Up（注册）或 Sign In（登录）",
            "按提示在浏览器中完成邮箱 / Google / GitHub 认证",
            "认证成功后自动返回 Cursor 客户端",
        ],
    )
    add_para(
        doc,
        "注意：Cursor 账号与电子邮箱绑定，现有账号的邮箱通常无法直接更改。企业用户可能使用 SSO（SAML），请按团队管理员指引操作。",
        size=10,
        color=RGBColor(0x66, 0x66, 0x66),
    )

    # ========== 下载 ==========
    add_heading_cn(doc, "三、下载官方安装包", 1)
    add_numbered(
        doc,
        [
            "打开官方下载页：https://cursor.com/download",
            "页面通常会自动识别你的操作系统；也可手动选择 Windows 或 macOS",
            "macOS 用户请确认芯片类型：Apple Silicon（M1/M2/M3/M4 等）或 Intel；不确定时可选择 Universal（通用）版本（若页面提供）",
            "点击对应下载按钮，等待安装包下载完成",
            "仅从官方站点下载，避免第三方镜像，降低被篡改风险",
        ],
    )
    add_image(doc, IMG / "01-download-page.png", width_cm=14.5, caption="图 2　官网下载页示意（Windows / macOS）")
    add_table(
        doc,
        ["平台", "典型文件名形态", "说明"],
        [
            ["Windows", "CursorSetup-*.exe 或类似名称", "运行安装向导；另有 User / System 安装类型"],
            ["macOS", "Cursor-*-*.dmg 或类似名称", "打开后将 Cursor 拖入「应用程序」"],
        ],
    )

    # ========== Windows ==========
    add_heading_cn(doc, "四、Windows 安装详细步骤", 1)
    add_heading_cn(doc, "4.1 运行安装程序", 2)
    add_numbered(
        doc,
        [
            "在「下载」文件夹中找到刚下载的 .exe 安装包，双击运行",
            "若出现 Windows SmartScreen「Windows 已保护你的电脑」提示：点击「更多信息」→「仍要运行」（请确认来源为 cursor.com）",
            "按安装向导提示操作：接受许可协议、确认安装选项",
            "安装类型（如有选择）：User（当前用户）通常无需管理员权限；System（系统级）安装到 Program Files，适合多用户共用",
            "建议勾选「添加到 PATH」或允许命令行使用 cursor 命令的相关选项（若向导提供）",
            "可按需勾选桌面快捷方式，然后点击 Install / 安装",
            "等待进度条完成，点击 Finish / 完成；部分版本会询问是否立即启动 Cursor",
        ],
    )
    add_image(doc, IMG / "03-windows-installer.png", width_cm=14.5, caption="图 3　Windows 安装向导示意")

    add_heading_cn(doc, "4.2 安装位置说明", 2)
    add_table(
        doc,
        ["类型", "默认路径"],
        [
            ["用户安装（User）", "%LOCALAPPDATA%\\Programs\\cursor\\（即 Cursor.exe）"],
            ["系统安装（System）", "%ProgramFiles%\\cursor\\"],
            ["用户数据 / 设置", "%APPDATA%\\Cursor\\"],
        ],
    )

    add_heading_cn(doc, "4.3 启动 Cursor（Windows）", 2)
    add_numbered(
        doc,
        [
            "从开始菜单搜索「Cursor」并打开，或双击桌面快捷方式",
            "首次启动进入欢迎 / 登录界面（见第六章）",
            "若启动白屏：可尝试以管理员身份运行，或重启后重试；仍失败请卸载后从官网重装",
        ],
    )

    # ========== Mac ==========
    add_heading_cn(doc, "五、macOS 安装详细步骤", 1)
    add_heading_cn(doc, "5.1 打开 DMG 并拖入应用程序", 2)
    add_numbered(
        doc,
        [
            "在「下载」中找到 .dmg 文件，双击打开磁盘镜像",
            "在打开的窗口中，将左侧的 Cursor 图标拖拽到右侧的 Applications（应用程序）文件夹",
            "等待复制完成（通常数秒到数十秒）",
            "在 Finder 侧边栏弹出已挂载的 Cursor 卷，右键「推出」；或关闭窗口后推出",
            "打开「启动台」或「应用程序」文件夹，点击 Cursor 启动",
        ],
    )
    add_image(doc, IMG / "04-mac-dmg-install.png", width_cm=14.5, caption="图 4　macOS：将 Cursor 拖入 Applications")

    add_heading_cn(doc, "5.2 首次打开与安全提示", 2)
    add_para(
        doc,
        "由于 Cursor 并非来自 Mac App Store，首次打开时，系统可能提示「来自互联网的应用」或需要在「隐私与安全性」中确认。",
    )
    add_numbered(
        doc,
        [
            "若弹出「是否打开」对话框：确认来源可信后点击「打开」",
            "若被拦截：打开「系统设置」→「隐私与安全性」，在相关提示处点击「仍要打开」",
            "也可在 Finder 中对 Cursor.app 右键 →「打开」，再确认一次",
            "安装位置一般为：/Applications/Cursor.app/",
        ],
    )
    add_heading_cn(doc, "5.3 Apple Silicon 与 Intel 说明", 2)
    add_bullets(
        doc,
        [
            "Apple Silicon（M 系列）：请优先下载 Apple Silicon / ARM64 版本，性能与后台索引更优",
            "Intel Mac：选择 Intel / x64 版本",
            "不确定芯片类型：点击苹果菜单 →「关于本机」查看芯片信息；或使用页面提供的 Universal 包",
        ],
    )
    add_heading_cn(doc, "5.4 macOS「已损坏」类提示（官方排障要点）", 2)
    add_para(
        doc,
        "若出现类似「Cursor 已损坏，无法打开」的提示，官方说明这通常与 macOS 校验有关，不一定是文件损坏：",
    )
    add_numbered(
        doc,
        [
            "完全退出 Cursor，在「活动监视器」中结束残留进程，等待约 1 分钟后重试",
            "仍不行：将 Cursor 移到废纸篓并清空，从 https://cursor.com/download 重新下载安装",
            "仍失败：重启 Mac 后再试",
        ],
    )

    # ========== 登录与首次使用 ==========
    add_heading_cn(doc, "六、首次启动、登录与初始化", 1)
    add_heading_cn(doc, "6.1 欢迎页与登录", 2)
    add_numbered(
        doc,
        [
            "启动 Cursor 后进入欢迎界面",
            "点击 Sign In（登录）或 Sign Up（注册）",
            "按提示使用邮箱、Google 或 GitHub 完成认证（可能跳转浏览器）",
            "认证成功后返回客户端，右上角通常可见头像 / 账号标识，表示已登录",
        ],
    )
    add_image(doc, IMG / "05-first-launch-signin.png", width_cm=14.5, caption="图 5　首次启动：欢迎与登录示意")

    add_heading_cn(doc, "6.2 可选：从 VS Code 导入", 2)
    add_para(
        doc,
        "若本机已安装 VS Code，首次启动向导可能提供「Import from VS Code」一类选项。可一键导入扩展、键位绑定与主题，缩短上手时间。不需要导入时可跳过，之后再在设置中处理。",
    )

    add_heading_cn(doc, "6.3 打开第一个项目", 2)
    add_numbered(
        doc,
        [
            "菜单：File（文件）→ Open Folder（打开文件夹）",
            "选择你的代码项目目录并确认",
            "等待项目索引（首次打开较大仓库时可能需要片刻）",
            "试用 AI：Windows 使用 Ctrl+I，macOS 使用 Cmd+I 打开 Agent；聊天面板常用 Ctrl+L / Cmd+L",
        ],
    )
    add_image(doc, IMG / "06-ready-to-code.png", width_cm=14.5, caption="图 6　登录成功后的编辑器示意")

    add_heading_cn(doc, "6.4 快速验证清单", 2)
    add_table(
        doc,
        ["检查项", "操作", "成功标志"],
        [
            ["客户端已安装", "从开始菜单 / 启动台打开 Cursor", "编辑器正常启动"],
            ["账号已登录", "查看右上角头像 / 账号入口", "显示已登录状态"],
            ["AI 可用", "打开 Chat 或 Agent 发一句简单问题", "有模型回复"],
            ["打开项目", "File → Open Folder", "资源管理器显示项目文件"],
        ],
    )

    # ========== 常见问题 ==========
    add_heading_cn(doc, "七、常见问题与排障", 1)
    add_heading_cn(doc, "7.1 下载或安装失败", 2)
    add_bullets(
        doc,
        [
            "确认从 https://cursor.com/download 下载，换浏览器或关闭下载加速插件后重试",
            "Windows：临时关闭干扰安装的安全软件，或以管理员身份运行安装包",
            "macOS：确认磁盘空间充足，并完全退出旧版后再覆盖安装到 Applications",
        ],
    )
    add_heading_cn(doc, "7.2 无法登录 / 网络异常", 2)
    add_bullets(
        doc,
        [
            "检查能否访问 cursor.com；公司网络可能拦截认证域名（如 *.cursor.sh）",
            "关闭 VPN 或换网络后重试；可改用 Google / GitHub 登录方式",
            "在 Cursor Settings → Network 中运行诊断；若代理不兼容 HTTP/2，可尝试 HTTP/1.1 兼容模式后重启",
        ],
    )
    add_heading_cn(doc, "7.3 更新客户端", 2)
    add_para(
        doc,
        "打开命令面板（Windows：Ctrl+Shift+P；macOS：Cmd+Shift+P），执行「Cursor: Attempt Update」，按提示重启完成更新。设置中可选择稳定版（Stable）或 Early Access 通道。",
    )

    # ========== 附录 ==========
    add_heading_cn(doc, "八、附录：常用链接与快捷键", 1)
    add_heading_cn(doc, "8.1 官方链接", 2)
    add_bullets(
        doc,
        [
            "下载：https://cursor.com/download",
            "安装帮助（中文）：https://cursor.com/cn/help/getting-started/install",
            "快速开始：https://cursor.com/docs/get-started/quickstart",
            "安装与启动排障：https://cursor.com/cn/help/troubleshooting/install-issues",
            "官网首页 / 注册：https://cursor.com",
        ],
    )
    add_heading_cn(doc, "8.2 入门快捷键", 2)
    add_table(
        doc,
        ["功能", "Windows", "macOS"],
        [
            ["Agent / Composer", "Ctrl + I", "Cmd + I"],
            ["Chat 聊天", "Ctrl + L", "Cmd + L"],
            ["行内编辑", "Ctrl + K", "Cmd + K"],
            ["命令面板", "Ctrl + Shift + P", "Cmd + Shift + P"],
            ["打开终端", "Ctrl + `", "Ctrl + `"],
        ],
    )

    add_para(
        doc,
        "—— 文档结束 ——",
        size=10,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        color=RGBColor(0x88, 0x88, 0x88),
        space_before=18,
    )
    add_para(
        doc,
        "配图为教学示意，界面细节以 Cursor 官方最新版本为准。",
        size=9,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        color=RGBColor(0x99, 0x99, 0x99),
    )

    doc.save(OUT)
    print(f"Wrote: {OUT}")


if __name__ == "__main__":
    build()
